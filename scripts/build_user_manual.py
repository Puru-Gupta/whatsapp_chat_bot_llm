from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_PATH = OUTPUT_DIR / "AQLI_WhatsApp_Chatbot_User_Manual.docx"
ICON_PATH = OUTPUT_DIR / "aqli_manual_icon.png"

FONT = "Calibri"
MONO = "Courier New"
INK = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEAL = "075E54"
MUTED = "65717C"
LIGHT_BLUE = "E8EEF5"
LIGHT_TEAL = "E9F5F2"
LIGHT_GRAY = "F2F4F7"
CAUTION = "7A5A00"
CAUTION_FILL = "FFF7DF"
RED = "9B1C1C"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, name=FONT, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]):
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError(f"Table widths must total {TABLE_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, separate, text, end])


def paragraph_shading(paragraph, fill: str, border_color: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border_color)
    p_bdr.append(left)
    p_pr.append(p_bdr)


def paragraph_bottom_rule(paragraph, color: str, size: str = "8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_numbering_definition(document: Document, abstract_id: int, num_id: int, bullet=False):
    numbering = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    if bullet:
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Calibri")
        fonts.set(qn("w:hAnsi"), "Calibri")
        r_pr.append(fonts)
        level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)


def add_numbering_instance(document: Document, abstract_id: int, num_id: int):
    numbering = document.part.numbering_part.element
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)


def apply_numbering(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def create_icon():
    image = Image.new("RGBA", (512, 512), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((16, 16, 496, 496), radius=110, fill="#075E54")
    draw.rounded_rectangle((132, 106, 380, 330), radius=70, fill="#E8FFF4")
    draw.polygon([(164, 302), (143, 400), (243, 330)], fill="#E8FFF4")
    draw.rounded_rectangle((190, 180, 326, 202), radius=11, fill="#128C7E")
    draw.rounded_rectangle((190, 238, 282, 260), radius=11, fill="#128C7E")
    image.save(ICON_PATH)


def configure_styles(document: Document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", 1)
    else:
        code = styles["Code Block"]
    code.font.name = MONO
    code._element.rPr.rFonts.set(qn("w:ascii"), MONO)
    code._element.rPr.rFonts.set(qn("w:hAnsi"), MONO)
    code.font.size = Pt(9.5)
    code.font.color.rgb = rgb(INK)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.15


def configure_page(document: Document):
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def set_running_furniture(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("AQLI WHATSAPP CHATBOT  |  USER MANUAL")
    set_run_font(run, size=8.5, color=MUTED, bold=True)
    paragraph_bottom_rule(p, "D7DEE6", "6")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run("AQLI User Manual  |  Page ")
    set_run_font(r, size=8.5, color=MUTED)
    add_page_field(p)


def add_body(document, text: str, bold_prefix: str | None = None, italic=False):
    p = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_bullet(document, text: str, bold_prefix: str | None = None):
    p = document.add_paragraph()
    apply_numbering(p, 21)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        set_run_font(p.add_run(text))
    return p


def add_step(document, text: str, bold_prefix: str | None = None, num_id: int = 22):
    p = document.add_paragraph()
    apply_numbering(p, num_id)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_callout(document, label: str, text: str, kind="info"):
    palette = {
        "info": (LIGHT_TEAL, TEAL),
        "note": (LIGHT_BLUE, BLUE),
        "caution": (CAUTION_FILL, CAUTION),
        "risk": ("FDECEC", RED),
    }
    fill, border = palette[kind]
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.2
    paragraph_shading(p, fill, border)
    set_run_font(p.add_run(f"{label}: "), bold=True, color=border)
    set_run_font(p.add_run(text), color=INK)
    return p


def add_code(document, text: str):
    p = document.add_paragraph(style="Code Block")
    paragraph_shading(p, LIGHT_GRAY, "C9D1DA")
    set_run_font(p.add_run(text), name=MONO, size=9.5, color=INK)
    return p


def add_table(document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        set_run_font(p.add_run(value), size=9.5, color=DARK_BLUE, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            set_run_font(p.add_run(value), size=9.5, color=INK)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    set_table_geometry(table, widths)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_heading(document, text: str, level=1, page_break=False):
    p = document.add_paragraph(text, style=f"Heading {level}")
    if page_break:
        p.paragraph_format.page_break_before = True
    return p


def add_cover(document: Document):
    section = document.sections[0]
    section.different_first_page_header_footer = True
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(35)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    icon = p.add_run().add_picture(str(ICON_PATH), width=Inches(1.12))
    icon._inline.docPr.set("descr", "AQLI WhatsApp chatbot logo")
    icon._inline.docPr.set("title", "AQLI chatbot")

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("AQLI WhatsApp Chatbot"), size=30, color=INK, bold=True)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    set_run_font(
        p.add_run("Professional User and Administrator Manual"),
        size=15,
        color=TEAL,
        bold=True,
    )

    rule = document.add_paragraph()
    rule.paragraph_format.space_after = Pt(26)
    paragraph_bottom_rule(rule, TEAL, "14")

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    set_run_font(
        p.add_run("Country and state/province AQLI analysis in WhatsApp"),
        size=12.5,
        color=DARK_BLUE,
    )

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(48)
    set_run_font(
        p.add_run("PM2.5 values | Life-expectancy impact | Rankings | Trends | Comparisons"),
        size=10.5,
        color=MUTED,
    )

    metadata = [
        ("Version", "1.0"),
        ("Issued", "4 July 2026"),
        ("WhatsApp", "+91 96547 01203"),
        ("Admin portal", "https://aqli-whatsapp-chatbot.vercel.app"),
    ]
    for label, value in metadata:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(f"{label}: "), size=10, color=MUTED, bold=True)
        set_run_font(p.add_run(value), size=10, color=MUTED)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    set_run_font(
        p.add_run("Prepared for AQLI chatbot users, analysts, and administrators"),
        size=9.5,
        color=MUTED,
        italic=True,
    )
    document.add_page_break()


def build_manual():
    create_icon()
    document = Document()
    configure_styles(document)
    configure_page(document)
    add_numbering_definition(document, 21, 21, bullet=True)
    add_numbering_definition(document, 22, 22, bullet=False)
    for num_id in range(23, 31):
        add_numbering_instance(document, 22, num_id)
    set_running_furniture(document.sections[0])
    add_cover(document)

    add_heading(document, "Document guide", 1)
    add_body(
        document,
        "This manual explains how to use the AQLI WhatsApp chatbot, how to interpret its answers, and how administrators maintain conversations and knowledge sources. It reflects the production behavior verified on 4 July 2026.",
    )
    add_callout(
        document,
        "Current data boundary",
        "The structured database contains country-level (GADM0) and state/province-level (GADM1) data through 2024. District-level (GADM2) data is not currently active.",
        "caution",
    )
    add_heading(document, "Contents", 2)
    contents = [
        "About the chatbot",
        "Quick start for WhatsApp users",
        "How to write a good question",
        "Supported question types",
        "Conversation memory and reset",
        "Understanding answers and references",
        "Data coverage, methodology, and limitations",
        "Troubleshooting for users",
        "Administrator guide",
        "Knowledge-base management",
        "Operational troubleshooting",
        "Testing checklist and prompt library",
        "Glossary",
    ]
    for item in contents:
        add_step(document, item)

    add_heading(document, "About the chatbot", 1, page_break=True)
    add_body(
        document,
        "The AQLI WhatsApp chatbot is a conversational interface for the Air Quality Life Index. It lets a user ask ordinary-language questions about long-term PM2.5 pollution and estimated life-expectancy impact without opening a spreadsheet, report, or analytical dashboard.",
    )
    add_heading(document, "What AQLI means", 2)
    add_body(
        document,
        "AQLI means Air Quality Life Index. It translates long-term exposure to fine particulate matter (PM2.5) into an estimated effect on life expectancy. The core relationship used by the chatbot is approximately 0.98 years of life-expectancy loss for each 10 µg/m³ increase in long-term PM2.5 above the selected benchmark.",
    )
    add_code(document, "Life loss per person = max(PM2.5 - benchmark, 0) x 0.098")
    add_body(
        document,
        "When the uploaded CSV already contains pre-calculated life-loss values, the chatbot uses those fields directly instead of recalculating them.",
    )
    add_heading(document, "Primary audiences", 2)
    for text in [
        "Policy and program teams that need quick, referenced AQLI facts.",
        "Researchers and analysts checking values, trends, rankings, and comparisons.",
        "Communications teams, journalists, educators, and public users seeking short explanations.",
        "Administrators who upload data and reports, monitor conversations, or take over a chat.",
    ]:
        add_bullet(document, text)

    add_heading(document, "Quick start for WhatsApp users", 1, page_break=True)
    add_heading(document, "Start a conversation", 2)
    add_step(document, "Open WhatsApp and start a chat with +91 96547 01203.", num_id=23)
    add_step(document, "Send a greeting such as Hi, Hello, or Good morning.", num_id=23)
    add_step(document, "Ask one focused question using a place, metric, and optional year or period.", num_id=23)
    add_step(document, "Read the answer, data-coverage notice, and reference line.", num_id=23)
    add_step(document, "Ask a follow-up or send reset to begin a new topic.", num_id=23)
    add_callout(
        document,
        "Fastest first test",
        "Send: What is India's PM2.5 in 2024? Then ask: Compare it with China.",
        "info",
    )
    add_heading(document, "Quick-reference examples", 2)
    add_table(
        document,
        ["Need", "Example question"],
        [
            ["Latest value", "What is India's latest PM2.5?"],
            ["State ranking", "Give top 5 most polluted states in India."],
            ["Threshold list", "List countries with PM2.5 above 30 in 2024."],
            ["Comparison", "Compare India and China from 2010 to 2020."],
            ["Mixed level", "Compare India and Uttar Pradesh for the last decade."],
            ["Trend", "Show the PM2.5 trend for Bihar."],
            ["Benchmark", "What is Jharkhand's life loss under the national standard?"],
            ["Methodology", "How does AQLI calculate life-expectancy loss?"],
        ],
        [2700, 6660],
    )

    add_heading(document, "How to write a good question", 1, page_break=True)
    add_body(
        document,
        "The chatbot handles natural wording, common misspellings, shuffled word order, and many comparison phrases. Accuracy is still highest when the question contains the following elements.",
    )
    add_table(
        document,
        ["Element", "What to provide", "Examples"],
        [
            ["Geography", "Country or state/province", "India; Bihar; Uttar Pradesh"],
            ["Metric", "PM2.5, life loss, population, standard", "PM2.5; life years lost"],
            ["Time", "A year, range, or relative period", "2024; 2010-2020; last decade"],
            ["Benchmark", "WHO, national, or custom target", "WHO; national std; target 15 µg/m³"],
            ["Operation", "Value, rank, compare, trend, or filter", "top 5; compare; above 30"],
        ],
        [1750, 3500, 4110],
    )
    add_heading(document, "Recommended question pattern", 2)
    add_code(document, "[Action] + [place/level] + [metric] + [year/period] + [benchmark]")
    add_body(document, "Example: Compare Bihar and Uttar Pradesh life loss from 2014 to 2024 using the national standard.")
    add_heading(document, "When the question is ambiguous", 2)
    add_body(
        document,
        "The chatbot asks one short clarification instead of guessing. For example, a request for top states without a country prompts for the country. A request for top 10 states every year may ask whether the user wants a period average or the single annual leader.",
    )
    add_code(document, "Bot: Do you want (1) the top 10 states averaged across the period, or (2) the single most polluted state for each year?\nUser: 2")

    add_heading(document, "Supported question types", 1, page_break=True)
    add_heading(document, "Direct values", 2)
    add_body(document, "Use direct-value questions for one location and one year. If no year is supplied, the latest available year is used. The current latest year is 2024.")
    for text in [
        "What is the PM2.5 level in Bihar?",
        "What was India's PM2.5 in 2020?",
        "What is the population of Nepal in the latest year?",
        "What is life-expectancy loss in Jharkhand under the national standard?",
    ]:
        add_bullet(document, text)

    add_heading(document, "Rankings", 2)
    add_body(document, "Rankings may be requested for countries or states/provinces. State rankings normally require a country unless the question explicitly asks for a worldwide state ranking.")
    for text in [
        "Top 10 countries by PM2.5 in 2024.",
        "Give top five most polluted states in Nepal.",
        "Which state has the highest life loss in India?",
        "Which countries improved the most over the last decade?",
    ]:
        add_bullet(document, text)
    add_callout(document, "Ranking rule", "PM2.5 and life loss are proportional for the same benchmark, so the bot normally shows both rather than asking which one to rank.", "note")

    add_heading(document, "Threshold filters", 2)
    add_body(document, "Threshold questions return matching places with units. The parser distinguishes strict, inclusive, and exact boundaries.")
    add_table(
        document,
        ["Meaning", "Recognized examples", "Behavior"],
        [
            ["Strictly above", "more than, greater than, above, over, exceeds, >", "Value must be greater than the threshold"],
            ["At or above", "at least, no less than, >=", "Threshold value is included"],
            ["Strictly below", "less than, lower than, below, under, <", "Value must be lower than the threshold"],
            ["At or below", "at most, no more than, <=", "Threshold value is included"],
            ["Equal", "equal to, exactly, same as, =", "Only exact matches are returned"],
        ],
        [1700, 3560, 4100],
    )
    add_body(document, "Examples: List states in India with PM2.5 below 30. List countries with national std more than 30. How many countries have PM2.5 at least 35?")

    add_heading(document, "Comparisons", 2)
    add_body(
        document,
        "The chatbot compares two countries, two states, or a country and a state. Country values are taken from GADM0; state/province values are taken from GADM1. Mixed-level answers include a level note so the user understands that the geographic units differ.",
    )
    add_code(document, "Compare India and China in 2024.\nCompare Bihar and Uttar Pradesh for the last decade.\nCompare India and Uttar Pradesh from 2014 to 2024.")
    add_body(document, "A period comparison shows the start value, end value, change, direction, final-year difference, and life-loss change for both places.")

    add_heading(document, "Trends and time periods", 2)
    add_body(
        document,
        "Trend questions show the first year, last year, PM2.5 change, life-loss change, and whether long-term exposure improved or worsened. Relative periods end at the latest available year. Use an explicit range when exact endpoints matter.",
    )
    for text in [
        "Show the trend for Madhesh.",
        "How did Delhi change from 2010 to 2024?",
        "Compare India and China over the past ten years.",
        "What was the highest PM2.5 in Uttar Pradesh during the last 20 years?",
    ]:
        add_bullet(document, text)

    add_heading(document, "Annual leaders and period averages", 2)
    add_body(document, "For multi-year rankings, the chatbot can calculate a period average or return the single leader for each year. If a top-N annual request could mean either, it asks the user to choose.")
    add_code(document, "Top 10 most polluted states in India between 2010 and 2024.\nGive the most polluted state in India for each year from 2010 to 2024.")

    add_heading(document, "Benchmarks", 2, page_break=True)
    add_table(
        document,
        ["Benchmark", "How to request it", "Data behavior"],
        [
            ["WHO guideline", "WHO, WHO guideline", "Default when no benchmark is specified; uses llpp_who when available"],
            ["National standard", "national standard, national std, nat std, national limit", "Uses llpp_nat/llpp_national when available"],
            ["Custom target", "target 10 µg/m³", "Calculates max(PM2.5 - target, 0) x 0.098"],
        ],
        [1800, 3000, 4560],
    )

    add_heading(document, "Total life-years calculation", 2)
    add_body(document, "When requested, the chatbot calculates an aggregate using the user-defined convention below.")
    add_code(document, "Total person-years lost = life loss per person x population")
    add_callout(document, "Interpret carefully", "This aggregate is a calculated communication measure. It should not be presented as an official individual prediction or medical diagnosis.", "caution")

    add_heading(document, "Reports and methodology", 2)
    add_body(document, "Use uploaded annual reports and methodology documents for definitions, policy context, interpretation, limitations, and report findings. Numeric rankings and comparisons continue to use the structured CSV data first.")
    for text in [
        "What is AQLI?",
        "Why does AQLI use PM2.5?",
        "How is life-expectancy loss calculated?",
        "What are the limitations of AQLI?",
        "Summarize the annual report's main findings.",
    ]:
        add_bullet(document, text)

    add_heading(document, "Conversation memory and reset", 1, page_break=True)
    add_heading(document, "Follow-up questions", 2)
    add_body(document, "The chatbot uses recent messages from the same conversation to interpret short follow-ups. It can retain the selected place, level, year, benchmark, comparison subject, or requested operation.")
    add_code(document, "User: Tell me about Uttar Pradesh.\nUser: Show its trend.\nUser: Compare it with Bihar.\nUser: Now use the national standard.")
    add_heading(document, "Start a clean topic", 2)
    add_body(document, "Send the single word reset when previous context is influencing a new question.")
    add_code(document, "reset")
    add_callout(document, "Important", "reset clears the active conversational context used for future answers. It does not delete stored conversation messages from Supabase.", "caution")
    add_heading(document, "When not to rely on memory", 2)
    for text in [
        "After changing from a country question to an unrelated state question.",
        "After a long sequence containing several benchmarks or time periods.",
        "When a reply mentions a place different from the one intended.",
        "Before formal testing, demonstrations, or screenshots.",
    ]:
        add_bullet(document, text)

    add_heading(document, "Understanding answers and references", 1, page_break=True)
    add_heading(document, "Standard answer structure", 2)
    add_step(document, "Answer first: the value, ranking, trend, or comparison requested.", num_id=24)
    add_step(document, "Short interpretation: included where it helps explain direction or significance.", num_id=24)
    add_step(document, "Suggested next question: a relevant follow-up, when available.", num_id=24)
    add_step(document, "Data coverage notice: country and state/province data only.", num_id=24)
    add_step(document, "Reference line: the source level, year/range, benchmark, and conversation context when used.", num_id=24)
    add_heading(document, "Reference labels", 2)
    add_table(
        document,
        ["Reference", "Meaning"],
        [
            ["GADM0 CSV", "Country-level structured data"],
            ["GADM1 CSV", "State/province-level structured data"],
            ["GADM0 CSV + GADM1 CSV", "Mixed country-to-state comparison using exact source levels"],
            ["AQLI Annual Report", "Narrative findings, context, and interpretation"],
            ["AQLI Methodology document", "Definitions, conversion relationship, assumptions, and limitations"],
            ["Previous conversation context + ...", "The current answer used a recent conversation selection"],
            ["Reference checked", "The bot checked the source but could not support the requested answer"],
        ],
        [2700, 6660],
    )
    add_heading(document, "Units", 2)
    for text in [
        "PM2.5 concentration: µg/m³.",
        "Life loss: years per person unless explicitly described as total person-years.",
        "Population: people.",
        "Change: µg/m³ or years, with direction stated separately.",
    ]:
        add_bullet(document, text)

    add_heading(document, "Data coverage, methodology, and limitations", 1, page_break=True)
    add_heading(document, "Current structured data coverage", 2)
    add_table(
        document,
        ["Level", "Status", "Use"],
        [
            ["GADM0 - Country", "Active", "Country values, rankings, filters, trends, and comparisons"],
            ["GADM1 - State/province", "Active", "State/province values, rankings, filters, trends, and comparisons"],
            ["GADM2 - District", "Not currently active", "District questions return a clear coverage limitation"],
        ],
        [2200, 2100, 5060],
    )
    add_heading(document, "Important limitations", 2)
    limitations = [
        ("Not live AQI", "The data represents annual long-term PM2.5, not current hourly air quality or an emergency alert."),
        ("Data freshness", "The current CSV series runs through 2024. New annual data must be uploaded before the bot can use it."),
        ("Grounded scope", "The chatbot does not invent unsupported values. Missing information produces a clear unavailable-data response."),
        ("Population estimates", "AQLI values describe estimated population-level effects, not an individual's medical outcome."),
        ("Geographic comparison", "Country-to-state comparisons are allowed but are explicitly labeled as different geographic levels."),
        ("Language", "The production parser is optimized for English. Other languages have not been fully validated."),
        ("Message format", "The webhook currently processes text messages only. Images, voice notes, documents, and other media are not answered automatically."),
        ("Platform dependency", "Replies depend on Meta WhatsApp Cloud API, Vercel, Supabase, and configured model/embedding services."),
        ("Conversation retention", "reset does not erase database records. Retention and deletion must be handled administratively."),
    ]
    for label, detail in limitations:
        add_bullet(document, f"{label}: {detail}", f"{label}: ")
    add_callout(document, "User-facing disclaimer", "Every automated reply states that current structured data coverage is country and state/province level only and that district-level data is not currently available.", "note")

    add_heading(document, "Troubleshooting for users", 1, page_break=True)
    add_table(
        document,
        ["Symptom", "Likely cause", "What to do"],
        [
            ["No reply", "Webhook/platform issue, human mode, or unsupported media message", "Send a text message. Wait briefly. Ask an administrator to check conversation mode and production logs."],
            ["Wrong place or topic", "Old conversation context", "Send reset, then ask the complete question again."],
            ["Bot asks for a country", "A state ranking did not include its country", "Reply with the country name, such as India or Nepal."],
            ["District unavailable", "No active GADM2 dataset", "Ask for country or state/province results."],
            ["No data for a year", "Requested year is outside the active CSV", "Use a supported year or ask for the latest available year."],
            ["Unexpected benchmark", "Benchmark was omitted or ambiguous", "State WHO, national standard, or a custom PM2.5 target explicitly."],
            ["Long answer", "The request asks for many places or years", "Request top 5, a shorter year range, or a specific location."],
        ],
        [2150, 3000, 4210],
    )
    add_heading(document, "A reliable retry sequence", 2)
    add_step(document, "Send reset.", num_id=25)
    add_step(document, "Name the geographic level and place.", num_id=25)
    add_step(document, "State the metric and exact year or range.", num_id=25)
    add_step(document, "State the benchmark if life loss is requested.", num_id=25)
    add_step(document, "If the problem continues, share the exact question and reply with the administrator.", num_id=25)

    add_heading(document, "Administrator guide", 1, page_break=True)
    add_heading(document, "Sign in", 2)
    add_step(document, "Open https://aqli-whatsapp-chatbot.vercel.app.", num_id=26)
    add_step(document, "Sign in with the Supabase administrator email and password.", num_id=26)
    add_step(document, "The application opens the conversation dashboard.", num_id=26)
    add_callout(document, "Security", "Never place production passwords, service-role keys, Meta tokens, or model API keys in this manual or in chat messages. Store them only in approved environment-variable systems.", "risk")

    add_heading(document, "Conversation dashboard", 2)
    for text in [
        "Search conversations by user name, phone number, or message text.",
        "Select a conversation to view its chronological message history.",
        "Unread counts reset when a conversation is opened.",
        "New messages and conversation updates appear through Supabase Realtime.",
        "Use the Knowledge base button to manage documents, CSVs, pasted text, and websites.",
    ]:
        add_bullet(document, text)
    add_heading(document, "AI mode and Human mode", 2)
    add_body(document, "Each conversation has a mode badge in the chat header.")
    add_table(
        document,
        ["Mode", "Behavior", "Use when"],
        [
            ["AI mode", "The chatbot automatically processes incoming text and sends a grounded reply.", "Normal automated service"],
            ["Human mode", "Incoming messages are stored, but automated replies stop. An administrator can reply manually.", "Escalations, sensitive questions, or manual support"],
        ],
        [1700, 4700, 2960],
    )
    add_step(document, "Open the required conversation.", num_id=27)
    add_step(document, "Select the AI mode/Human mode badge to switch modes.", num_id=27)
    add_step(document, "In Human mode, type a response in the message field and send it.", num_id=27)
    add_step(document, "Switch back to AI mode when automated replies should resume.", num_id=27)

    add_heading(document, "Knowledge-base management", 1, page_break=True)
    add_heading(document, "Open the knowledge manager", 2)
    add_body(document, "From the conversation dashboard, select Knowledge base. The manager supports pasted text, file upload, and website import.")
    add_heading(document, "Supported source types", 2)
    add_table(
        document,
        ["Source", "Limits", "Recommended use"],
        [
            ["Pasted text", "Up to 2,000,000 characters", "FAQs, policy notes, corrections, short methodology text"],
            ["PDF", "25 MB maximum", "Annual reports and methodology reports"],
            ["DOCX", "25 MB maximum", "Structured narrative documents"],
            ["CSV", "25 MB maximum", "GADM0/GADM1 structured numeric datasets"],
            ["TXT/Markdown", "25 MB maximum", "Plain-text references and prepared notes"],
            ["Website", "Single page or up to 12 indexed site pages", "Public AQLI pages and selected supporting content"],
        ],
        [1800, 2500, 5060],
    )

    add_heading(document, "Upload a file", 2)
    add_step(document, "Enter a clear source title, including level and data year where relevant.", num_id=28)
    add_step(document, "Select Upload file and choose a PDF, DOCX, CSV, TXT, or Markdown file.", num_id=28)
    add_step(document, "Select Import file and wait for the chunk count confirmation.", num_id=28)
    add_step(document, "Confirm that the source status is Active.", num_id=28)
    add_step(document, "Run a representative WhatsApp question and verify the reference line.", num_id=28)

    add_heading(document, "CSV requirements", 2)
    add_body(document, "The current wide-year schema uses one row per geographic unit with yearly PM2.5 and life-loss columns. Common fields are shown below.")
    add_table(
        document,
        ["Field family", "Meaning"],
        [
            ["country", "Country name"],
            ["name_1", "State or province name"],
            ["name_2", "District name; not active in the current production dataset"],
            ["region, continent, iso_alpha3", "Geographic classification"],
            ["population", "Population used for weighted averages or aggregate calculations"],
            ["whostandard, natstandard", "WHO and national PM2.5 benchmark fields"],
            ["pm1998 ... pm2024", "Annual average PM2.5 values"],
            ["llpp_who_1998 ... llpp_who_2024", "Life loss per person relative to WHO"],
            ["llpp_nat_1998 ... llpp_nat_2024", "Life loss per person relative to the national standard"],
        ],
        [3000, 6360],
    )
    add_callout(document, "Replacement behavior", "Uploading a newer CSV for the same GADM level automatically deactivates older active CSV sources at that level. Verify the level classification and active status after every annual replacement.", "info")

    add_heading(document, "Import a website", 2)
    add_step(document, "Enter a descriptive title and a public website URL.", num_id=29)
    add_step(document, "Choose Single page for one URL or Site for a limited same-site crawl.", num_id=29)
    add_step(document, "Select Import website. Site mode indexes up to 12 pages per import.", num_id=29)
    add_step(document, "Review pages indexed, warnings, and generated chunk count.", num_id=29)
    add_step(document, "Test a question whose answer appears clearly on the imported page.", num_id=29)
    add_callout(document, "Website limitation", "Dynamic, login-protected, blocked, oversized, or script-only pages may not provide usable text. Import a clean PDF or pasted text when necessary.", "caution")

    add_heading(document, "Manage existing sources", 2)
    for text in [
        "Deactivate a source to stop using it without deleting it.",
        "Activate a source to make it available to future answers.",
        "Reprocess a source when extraction or indexing must be rebuilt.",
        "Delete a source to remove it and its associated knowledge chunks.",
        "Keep only the intended current GADM0 and GADM1 CSVs active.",
    ]:
        add_bullet(document, text)

    add_heading(document, "Operational troubleshooting", 1, page_break=True)
    add_table(
        document,
        ["Problem", "Administrator checks", "Resolution"],
        [
            ["No WhatsApp reply", "Check Vercel POST /api/webhook logs; conversation mode; Meta account status; messages webhook subscription", "Restore Meta access/subscription, correct credentials, or return conversation to AI mode"],
            ["Webhook verifies but replies fail", "Check outbound Meta error and configured Phone Number ID", "Update WhatsApp credentials/IDs in Vercel and redeploy"],
            ["Wrong old data", "Check active knowledge sources and CSV level classification", "Upload the replacement CSV; deactivate or delete stale sources"],
            ["Correct question uses old context", "Review recent messages and resolved follow-up sequence", "Ask user to send reset; reproduce question as standalone input"],
            ["PDF/CSV upload fails", "Check format, 25 MB limit, extractable text, and route logs", "Clean or split the source; retry with a supported file"],
            ["Website content not used", "Check source status, indexed pages, warnings, and visible page text", "Reprocess, import single page, or paste authoritative text"],
            ["Slow response", "Check grounded-answer elapsed time, model fallback, source size, and Vercel duration", "Use a narrower question; review retrieval/model latency and source quality"],
        ],
        [1900, 3970, 3490],
    )
    add_heading(document, "Environment configuration categories", 2)
    add_body(document, "Production requires correctly configured environment variables in Vercel. Do not record values in operational documents.")
    for text in [
        "Supabase: project URL, publishable key, service-role key.",
        "WhatsApp: verify token, access token, Phone Number ID, Business Account ID.",
        "Model and embeddings: OpenRouter/model setting and embedding provider credentials as applicable.",
        "Application: APP_URL and NEXT_PUBLIC_APP_URL.",
    ]:
        add_bullet(document, text)
    add_heading(document, "Production health check", 2, page_break=True)
    add_step(document, "Confirm the latest Vercel deployment status is Ready.", num_id=30)
    add_step(document, "Verify GET /api/webhook with the configured verify token returns HTTP 200.", num_id=30)
    add_step(document, "Send a real inbound WhatsApp text and confirm POST /api/webhook appears in logs.", num_id=30)
    add_step(document, "Confirm a grounded answer is stored in Supabase and delivered by WhatsApp.", num_id=30)
    add_step(document, "Check recent error logs for Meta send failures, timeouts, and source-processing errors.", num_id=30)

    add_heading(document, "Testing checklist and prompt library", 1)
    add_heading(document, "Release acceptance checklist", 2)
    checklist = [
        "Greeting and thanks produce natural responses.",
        "reset clears active context but preserves message history.",
        "Country direct value uses GADM0 and latest year by default.",
        "State direct value uses GADM1.",
        "Top-state request without country asks for the country.",
        "Strict, inclusive, and equal threshold operators return correct boundaries.",
        "Explicit and relative period comparisons show both places.",
        "Country-to-state comparison cites GADM0 + GADM1 and includes a level note.",
        "WHO, national std, nat std, and custom target select the intended benchmark.",
        "Common misspellings and joined place names are handled or clarified.",
        "District requests state that GADM2 is unavailable.",
        "Every reply includes a reference and data-coverage notice.",
        "Human mode stops automated replies; AI mode resumes them.",
    ]
    for item in checklist:
        add_bullet(document, item)

    add_heading(document, "Recommended test prompts", 2, page_break=True)
    prompts = [
        "What is India's PM2.5 in 2024?",
        "Give top five states in Nepal by PM2.5.",
        "List countries with national std more than 30.",
        "List states in India with PM2.5 at or below 30.",
        "Compare India and China between 2010 and 2020.",
        "Compare India and Uttar Pradesh for the last decade.",
        "Compare Bihar and Uttar Pradesh life loss under the national standard.",
        "What is national-standard life loss for Jharkhand and Uttar Pradesh for the last two years?",
        "Show the trend for Madhesh.",
        "Give the most polluted state in India for each year from 2010 to 2024.",
        "What is AQLI and how is life loss calculated?",
        "Top 10 polluted districts in India.",
    ]
    for prompt in prompts:
        add_code(document, prompt)

    add_heading(document, "Glossary", 1, page_break=True)
    add_table(
        document,
        ["Term", "Definition"],
        [
            ["AQLI", "Air Quality Life Index; translates long-term PM2.5 exposure into estimated life-expectancy impact"],
            ["PM2.5", "Fine particulate matter with aerodynamic diameter of 2.5 micrometers or smaller"],
            ["GADM0", "Country-level geographic dataset"],
            ["GADM1", "State/province-level geographic dataset"],
            ["GADM2", "District-level geographic dataset; not currently active"],
            ["WHO benchmark", "World Health Organization PM2.5 guideline field used as the default benchmark"],
            ["National standard", "Country-specific PM2.5 standard from the active dataset"],
            ["llpp", "Life loss per person"],
            ["Period average", "Average value across the selected years for each place"],
            ["Annual leader", "The highest or lowest place for each individual year"],
            ["AI mode", "Conversation mode in which automated replies are enabled"],
            ["Human mode", "Conversation mode in which automatic replies stop and an administrator responds"],
            ["Knowledge source", "An active or inactive report, website, text, or CSV available to the chatbot"],
            ["Grounded answer", "An answer supported by the active CSV, report, methodology document, or recent conversation context"],
        ],
        [2200, 7160],
    )

    add_heading(document, "Support and document control", 1)
    add_body(document, "Production application: https://aqli-whatsapp-chatbot.vercel.app")
    add_body(document, "WhatsApp chatbot number: +91 96547 01203")
    add_body(document, "Document version: 1.0 | Issued: 4 July 2026")
    add_body(document, "Update this manual whenever data coverage, benchmark rules, supported message types, administrator workflows, or production platform configuration changes.")

    configure_page(document)
    document.core_properties.title = "AQLI WhatsApp Chatbot User Manual"
    document.core_properties.subject = "Professional user and administrator guide"
    document.core_properties.author = "AQLI"
    document.core_properties.keywords = "AQLI, WhatsApp, chatbot, PM2.5, user manual"
    document.core_properties.comments = "Production behavior verified 4 July 2026"
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_manual()
